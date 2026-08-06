import io
import tempfile

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase, override_settings

from openpyxl import Workbook, load_workbook

from . import services


def _build_xlsx(sheet_name, headers, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(headers)
    for row in rows:
        ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    wb.close()
    return buffer.getvalue()


class HeaderMapTests(SimpleTestCase):

    def _build(self, name, sheet_name, headers, rows, token):
        media = tempfile.mkdtemp()
        with override_settings(MEDIA_ROOT=media):
            content = _build_xlsx(sheet_name, headers, rows)
            uploaded = SimpleUploadedFile(
                name, content,
                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            )
            return services.detect_and_build_dir(uploaded, token)

    def _read_slot(self, base_dir, slot_rel_path, sheet):
        wb = load_workbook(base_dir / slot_rel_path, data_only=True)
        try:
            return list(wb[sheet].iter_rows(values_only=True))
        finally:
            wb.close()

    def test_reconoce_docentes_por_encabezados_y_reordena_columnas(self):
        base_dir, labels = self._build(
            'docentes_mio.xlsx', 'Mi hoja',
            ['Correo', 'Nombres completos', 'Cédula', 'Celular',
             'Modalidad', 'Dedicación', 'Tipo de sangre'],
            [['juan@uleam.edu.ec', 'JUAN PEREZ', '1301234567', '0991234567',
              'CONTRATO', 'TIEMPO COMPLETO', 'O+']],
            'tok_doc',
        )
        self.assertIn('Catálogo de docentes', labels)
        rows = self._read_slot(base_dir, 'Docentes.xlsx', 'MDOCENTES')
        data = rows[1]
        self.assertEqual(data[0], '1301234567')      # cedula
        self.assertEqual(data[1], 'JUAN PEREZ')      # nombres
        self.assertEqual(data[3], 'juan@uleam.edu.ec')
        self.assertEqual(data[4], '0991234567')
        self.assertEqual(data[6], 'TIEMPO COMPLETO')  # dedicacion
        self.assertEqual(data[7], 'CONTRATO')         # modalidad
        self.assertEqual(data[8], 'O+')
        self.assertIsNone(data[5])

    def test_reconoce_planificacion_por_encabezados(self):
        base_dir, labels = self._build(
            'planificacion_mia.xlsx', 'Planificación',
            ['Docente', 'Paralelo', 'Horas', 'Asignatura', 'Carrera',
             'Nivel', 'Campo', 'Cédula docente'],
            [['ANA LOZANO', 'A', 4, 'CONTABILIDAD', 'ADMINISTRACION', 5,
              'CONTABILIDAD', '1312345678']],
            'tok_plan',
        )
        self.assertIn('Planificación / asignaciones', labels)
        rows = self._read_slot(base_dir, 'FCACC-PLANIFICACION.xlsx', 'ASIGNACION')
        data = rows[1]
        self.assertEqual(data[1], 'ADMINISTRACION')   # carrera
        self.assertEqual(data[3], 5)                  # nivel
        self.assertEqual(data[4], 'A')                # paralelo
        self.assertEqual(data[5], 'CONTABILIDAD')     # asignatura
        self.assertEqual(data[7], 'CONTABILIDAD')     # campo
        self.assertEqual(data[11], 4)                 # horas
        self.assertEqual(data[12], 'ANA LOZANO')      # docente
        self.assertEqual(data[22], '1312345678')      # cedula

    def test_archivo_sin_hojas_reconocidas_lanza_error(self):
        media = tempfile.mkdtemp()
        with override_settings(MEDIA_ROOT=media):
            content = _build_xlsx('Data', ['X', 'Y', 'Z'], [[1, 2, 3]])
            uploaded = SimpleUploadedFile('otro.xlsx', content)
            with self.assertRaises(services.ArchivoNoReconocido):
                services.detect_and_build_dir(uploaded, 'tok_no')

    def test_archivo_no_excel_lanza_error_amigable(self):
        media = tempfile.mkdtemp()
        with override_settings(MEDIA_ROOT=media):
            uploaded = SimpleUploadedFile('basura.xlsx', b'not an excel at all')
            with self.assertRaises(services.ArchivoNoReconocido):
                services.detect_and_build_dir(uploaded, 'tok_bad')
