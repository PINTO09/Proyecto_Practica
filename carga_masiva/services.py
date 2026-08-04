"""Motor de la carga masiva vía Excel: guarda los archivos subidos con la
estructura que espera ``import_complete_fcacc``, ejecuta ese comando (ya
probado) dentro de una transacción controlada para calcular una vista previa
sin persistir nada, y aplica los cambios reales + auditoría al confirmar.

No reimplementa ninguna regla de negocio del importador: solo lo envuelve.
"""
import io
import shutil
import uuid
from pathlib import Path

from django.conf import settings
from django.db import transaction

from catalogos.models import (
    CatalogoCampoConocimiento, CatalogoCarrera, CatalogoPeriodoAcademico,
    CatalogoTituloPosgrado,
)
from curriculo.models import (
    CurriculoAsignatura, CurriculoAsignaturaCampo, RelacionPosgradoCampo,
)
from docentes.models import DocenteCampoAfinidad, DocenteFcacc, DocenteTituloAcademico
from planificacion.management.commands.import_complete_fcacc import (
    CAREER_ALIASES, F4_ACTIVITY_CAREER_CODES, Command as ImportCommand,
    _canonical_name, _clean_text, _fit_code, _normalize_cedula,
    _normalize_text, _valid_level,
)
from planificacion.models import (
    CatalogoActividadComplementaria, PlanificacionActividadDocente,
    PlanificacionAsignacionDocente, PlanificacionDemandaAcademica,
)

# Modelos que toca el importador, en el mismo orden en que los procesa.
TRACKED_MODELS = [
    (CatalogoCarrera, 'Carreras'),
    (CatalogoCampoConocimiento, 'Campos de conocimiento'),
    (CatalogoTituloPosgrado, 'Títulos de posgrado'),
    (CatalogoPeriodoAcademico, 'Períodos académicos'),
    (CurriculoAsignatura, 'Asignaturas'),
    (CurriculoAsignaturaCampo, 'Relación asignatura-campo'),
    (RelacionPosgradoCampo, 'Relación posgrado-campo'),
    (DocenteFcacc, 'Docentes'),
    (DocenteTituloAcademico, 'Títulos académicos de docentes'),
    (DocenteCampoAfinidad, 'Afinidad docente-campo'),
    (PlanificacionDemandaAcademica, 'Demanda académica'),
    (PlanificacionAsignacionDocente, 'Asignaciones docentes'),
    (CatalogoActividadComplementaria, 'Catálogo de actividades complementarias'),
    (PlanificacionActividadDocente, 'Actividades asignadas a docentes'),
]

# Campos representativos para mostrar cada fila en la vista previa sin volcar
# todas las columnas crudas.
_DISPLAY_FIELDS = {
    CatalogoCarrera: ('codigo_carrera', 'nombre_carrera'),
    CatalogoCampoConocimiento: ('codigo_campo', 'nombre_campo_conocimiento'),
    CatalogoTituloPosgrado: ('codigo_posgrado', 'nombre_titulo_posgrado'),
    CatalogoPeriodoAcademico: ('codigo_periodo', 'nombre_periodo'),
    CurriculoAsignatura: ('codigo_asignatura', 'nombre_asignatura'),
    CurriculoAsignaturaCampo: ('id_asignatura_id', 'id_campo_id'),
    RelacionPosgradoCampo: ('id_posgrado_id', 'id_campo_id'),
    DocenteFcacc: ('cedula_docente', 'nombres_completos'),
    DocenteTituloAcademico: ('id_docente_id', 'nombre_titulo'),
    DocenteCampoAfinidad: ('id_docente_id', 'id_campo_id'),
    PlanificacionDemandaAcademica: ('id_asignatura_id', 'id_carrera_id', 'numero_paralelos'),
    PlanificacionAsignacionDocente: ('id_asignatura_id', 'id_docente_id', 'paralelo_asignado'),
    CatalogoActividadComplementaria: ('codigo_actividad', 'nombre_actividad'),
    PlanificacionActividadDocente: ('id_docente_id', 'id_actividad_id', 'horas_asignadas'),
}


# ─── Exclusión de filas puntuales ──────────────────────────────────────────
# Para "no incluir esta fila", se necesita encontrar de vuelta la fila cruda
# del Excel que la produjo y quitarla antes de volver a correr el importador
# (así todas las validaciones del comando se siguen aplicando igual). Cada
# builder recibe los datos ya resueltos del registro (mismo formato que
# _model_to_dict) y devuelve [(nombre_hoja, predicado_de_fila)]. Se reusan
# las mismas funciones de normalización que ya usa el importador para
# garantizar que el "match" inverso sea exactamente igual al "match" original.

def _excl_carrera(data):
    codigo = data.get('codigo_carrera')
    return [('MAE_CARRERA', lambda row: bool(row) and row[0] is not None and _clean_text(row[0]) == codigo)]


def _excl_campo(data):
    codigo = data.get('codigo_campo')
    return [('MAE_CONOCIMIENTO', lambda row: bool(row) and row[0] is not None and _clean_text(row[0]) == codigo)]


def _excl_posgrado(data):
    nombre = _normalize_text(data.get('nombre_titulo_posgrado') or '')
    return [('MAESTRIA', lambda row: len(row or ()) > 1 and row[1] is not None and _normalize_text(row[1]) == nombre)]


def _excl_asignatura(data):
    codigo = data.get('codigo_asignatura')

    def pred(row):
        if not row or row[0] is None:
            return False
        return _fit_code(_clean_text(row[0]), 'ASG-') == codigo

    return [('MAE_ASIGNATURA', pred)]


def _excl_docente(data):
    cedula = data.get('cedula_docente')
    return [('MDOCENTES', lambda row: bool(row) and row[0] is not None and _normalize_cedula(row[0]) == cedula)]


def _resolve_asignacion_identity(data):
    """Resuelve (carrera canonica, nivel, asignatura canonica) para una fila
    de demanda/asignacion a partir de sus FK, igual que hace el importador."""
    asignatura = CurriculoAsignatura.objects.select_related('id_carrera').get(pk=data['id_asignatura_id'])
    carrera_canonica = _canonical_name(asignatura.id_carrera.nombre_carrera)
    carrera_canonica = CAREER_ALIASES.get(carrera_canonica, carrera_canonica)
    return carrera_canonica, asignatura.nivel_semestre, _canonical_name(asignatura.nombre_asignatura)


def _match_asignacion_row(row, carrera_canonica, nivel, subject_canonica):
    if len(row) <= 12:
        return False
    carrera_nombre = _clean_text(row[1])
    row_nivel = _valid_level(row[3])
    subject_nombre = _clean_text(row[5])
    if not carrera_nombre or not subject_nombre or row_nivel is None:
        return False
    row_carrera = _canonical_name(carrera_nombre)
    row_carrera = CAREER_ALIASES.get(row_carrera, row_carrera)
    row_subject = _canonical_name(subject_nombre)
    return row_carrera == carrera_canonica and row_nivel == nivel and row_subject == subject_canonica


def _excl_demanda(data):
    carrera_canonica, nivel, subject_canonica = _resolve_asignacion_identity(data)

    def pred(row):
        return _match_asignacion_row(row, carrera_canonica, nivel, subject_canonica)

    return [('ASIGNACION', pred)]


def _excl_asignacion_docente(data):
    carrera_canonica, nivel, subject_canonica = _resolve_asignacion_identity(data)
    paralelo = (data.get('paralelo_asignado') or 'A').strip()[:3].upper()
    docente = DocenteFcacc.objects.get(pk=data['id_docente_id'])
    teacher_norm = _normalize_text(docente.nombres_completos)

    def pred(row):
        if not _match_asignacion_row(row, carrera_canonica, nivel, subject_canonica):
            return False
        row_paralelo = (_clean_text(row[4]) or 'A')[:3].upper()
        if row_paralelo != paralelo:
            return False
        row_teacher = _clean_text(row[12]) if len(row) > 12 else ''
        return _normalize_text(row_teacher) == teacher_norm

    return [('ASIGNACION', pred)]


def _excl_actividad_docente(data):
    actividad = CatalogoActividadComplementaria.objects.get(pk=data['id_actividad_id'])
    docente = DocenteFcacc.objects.get(pk=data['id_docente_id'])
    codigo_actividad = actividad.codigo_actividad
    teacher_norm = _normalize_text(docente.nombres_completos)

    def pred(row):
        if len(row) <= 12:
            return False
        career_code = _clean_text(row[2]) if len(row) > 2 else ''
        activity_field = _normalize_text(row[7]) if len(row) > 7 else ''
        if career_code not in F4_ACTIVITY_CAREER_CODES and activity_field != 'ACTIVIDAD':
            return False
        row_teacher = _clean_text(row[12])
        if _normalize_text(row_teacher) != teacher_norm:
            return False
        activity_name = _clean_text(row[5])
        source_code = _clean_text(row[6]) or activity_name
        return _fit_code(source_code, 'ACT-') == codigo_actividad

    return [('ASIGNACION', pred)]


# Solo se ofrece excluir filas de las entidades donde se puede reconstruir
# con precisión la fila de origen. Las tablas de relación (asignatura-campo,
# posgrado-campo, afinidad, titulos) se muestran de solo lectura en la vista
# previa para no arriesgar una exclusion imprecisa.
EXCLUSION_BUILDERS = {
    CatalogoCarrera: _excl_carrera,
    CatalogoCampoConocimiento: _excl_campo,
    CatalogoTituloPosgrado: _excl_posgrado,
    CurriculoAsignatura: _excl_asignatura,
    DocenteFcacc: _excl_docente,
    PlanificacionDemandaAcademica: _excl_demanda,
    PlanificacionAsignacionDocente: _excl_asignacion_docente,
    PlanificacionActividadDocente: _excl_actividad_docente,
}


def exclude_key(model, pk):
    return f'{model._meta.label}:{pk}'

# Cada "slot" es uno de los archivos que espera import_complete_fcacc.
# (ruta relativa, hojas que debe contener su archivo "stub" vacío para que
# el comando no falle por hoja faltante, etiqueta)
SLOT_SPECS = {
    'principal': ('FCACC-PLANIFICACION.xlsx', ['ASIGNACION', 'MAE_ASIGNATURA', 'EDU_DOCENTE'], 'Planificación / asignaciones'),
    'revision_docentes': ('revision/REVISIONDocentes.xlsx', ['DET_DOCENTE'], 'Afinidad de docentes'),
    'nivel_docente': ('NIVEL_DOCENTE.xlsx', ['EDU_DOCENTE'], 'Nivel académico de docentes'),
    'docentes': ('Docentes.xlsx', ['MDOCENTES'], 'Catálogo de docentes'),
    'maestria': ('Maestria.xlsx', ['MAESTRIA'], 'Catálogo de maestrías'),
    'detallado_maestria': ('DETALLADO_MAESTRIA.xlsx', ['MAESTRIA_DETALLADO'], 'Detalle maestría-campo'),
    'detallado': ('DETALLADO.xlsx', ['MAE_CONOCIMIENTO'], 'Campos de conocimiento'),
    'detallado_asignatura': ('DETALLADO_ASIGNATURA.xlsx', ['MAE_CARRERA', 'MAE_ASIGNATURA', 'DET_ASIG'], 'Carreras / asignaturas'),
}

# Nombre de hoja detectado en el archivo subido -> slot(s) que debe llenar.
# Una hoja puede alimentar más de un slot (p.ej. EDU_DOCENTE aparece tanto en
# el libro principal como en NIVEL_DOCENTE.xlsx).
SHEET_TO_SLOTS = {
    'ASIGNACION': ['principal'],
    'EDU_DOCENTE': ['principal', 'nivel_docente'],
    'MAE_ASIGNATURA': ['principal', 'detallado_asignatura'],
    'DET_DOCENTE': ['revision_docentes'],
    'MDOCENTES': ['docentes'],
    'MAESTRIA': ['maestria'],
    'MAESTRIA_DETALLADO': ['detallado_maestria'],
    'MAE_CONOCIMIENTO': ['detallado'],
    'MAE_CARRERA': ['detallado_asignatura'],
    'DET_ASIG': ['detallado_asignatura'],
}


class _PreviewRollback(Exception):
    """Señal interna para forzar el rollback de la vista previa."""


class ArchivoNoReconocido(ValueError):
    """El Excel subido no tiene ninguna hoja que reconozcamos."""


def temp_dir(token):
    return Path(settings.MEDIA_ROOT) / 'carga_masiva' / 'tmp' / token


def _write_merged(dest_path, source_wb, sheet_names):
    """Crea el Excel de un slot con exactamente las hojas que ese rol
    necesita: copia los datos reales del archivo subido para las hojas que sí
    trae, y deja una hoja vacía (cero filas) para las que no trae. Así el
    importador nunca falla por "hoja faltante", sin importar qué combinación
    de hojas haya subido el usuario."""
    from openpyxl import Workbook
    out = Workbook()
    default_sheet = out.active
    for i, name in enumerate(sheet_names):
        ws = default_sheet if i == 0 else out.create_sheet()
        ws.title = name
        if name in source_wb.sheetnames:
            for row in source_wb[name].iter_rows(values_only=True):
                ws.append(row)
    out.save(dest_path)


def detect_and_build_dir(uploaded_file, token):
    """Analiza las hojas del Excel subido, coloca ese archivo en los slots
    que le correspondan y rellena el resto con archivos vacíos (mismas hojas,
    cero filas) para que import_complete_fcacc solo procese la información
    de este archivo. Devuelve (base_dir, etiquetas_detectadas).

    Levanta ArchivoNoReconocido si ninguna hoja del archivo coincide con lo
    que espera el importador.
    """
    from openpyxl import load_workbook

    base_dir = temp_dir(token)
    base_dir.mkdir(parents=True, exist_ok=True)

    upload_path = base_dir / '_subido.xlsx'
    with open(upload_path, 'wb') as f:
        for chunk in uploaded_file.chunks():
            f.write(chunk)

    # data_only=True para que las celdas con formulas traigan su ultimo valor
    # calculado en vez del texto de la formula.
    wb = load_workbook(upload_path, data_only=True)
    sheet_names = set(wb.sheetnames)

    matched_slots = set()
    for sheet in sheet_names:
        matched_slots.update(SHEET_TO_SLOTS.get(sheet, []))

    if not matched_slots:
        wb.close()
        upload_path.unlink(missing_ok=True)
        raise ArchivoNoReconocido(
            'No reconocemos ninguna hoja de este archivo para la carga masiva. '
            'Hojas encontradas: ' + ', '.join(sorted(sheet_names))
        )

    # Cada slot se arma siempre por fusion (datos reales donde el archivo
    # subido tenga esa hoja exacta, vacio donde no) -- asi un archivo que
    # solo trae ALGUNAS de las hojas de un slot (p.ej. solo ASIGNACION, sin
    # EDU_DOCENTE) nunca deja una hoja faltante que rompa al importador.
    for relative_path, stub_sheets, _label in SLOT_SPECS.values():
        dest = base_dir / relative_path
        dest.parent.mkdir(parents=True, exist_ok=True)
        _write_merged(dest, wb, stub_sheets)

    wb.close()
    upload_path.unlink(missing_ok=True)
    detected_labels = sorted(SLOT_SPECS[slot][2] for slot in matched_slots)
    return base_dir, detected_labels


def apply_exclusions(base_dir, excluded_pairs):
    """Reescribe en base_dir las hojas necesarias para quitar las filas
    crudas correspondientes a excluded_pairs (lista de (modelo, data_dict))
    antes de que se vuelva a correr el importador. Si un tipo de dato no
    tiene builder de exclusion (relaciones), se ignora silenciosamente."""
    if not excluded_pairs:
        return

    predicates_by_sheet = {}
    for model, data in excluded_pairs:
        builder = EXCLUSION_BUILDERS.get(model)
        if not builder:
            continue
        for sheet_name, pred in builder(data):
            predicates_by_sheet.setdefault(sheet_name, []).append(pred)

    if not predicates_by_sheet:
        return

    from openpyxl import load_workbook

    for relative_path, stub_sheets, _label in SLOT_SPECS.values():
        relevant_sheets = [s for s in stub_sheets if s in predicates_by_sheet]
        if not relevant_sheets:
            continue
        path = base_dir / relative_path
        if not path.exists():
            continue
        wb = load_workbook(path, data_only=True)
        changed = False
        for sheet_name in relevant_sheets:
            if sheet_name not in wb.sheetnames:
                continue
            rows = list(wb[sheet_name].iter_rows(values_only=True))
            if not rows:
                continue
            header, data_rows = rows[0], rows[1:]
            preds = predicates_by_sheet[sheet_name]
            kept = [r for r in data_rows if not any(p(r) for p in preds)]
            if len(kept) == len(data_rows):
                continue
            changed = True
            del wb[sheet_name]
            new_ws = wb.create_sheet(sheet_name)
            new_ws.append(header)
            for r in kept:
                new_ws.append(r)
        if changed:
            wb.save(path)
        wb.close()


def excludable_items_from_diff(diff):
    """Extrae de un diff ya calculado {exclude_key: data_dict} para las
    filas nuevas/modificadas de entidades excluibles. Los datos de las
    filas "nuevas" solo existen en este momento (la vista previa se revierte
    siempre), por eso se guardan en sesión en vez de volver a consultar la
    BD al confirmar."""
    result = {}
    for entity in diff.values():
        if not entity.get('excludable'):
            continue
        for item in entity['new']:
            if item['exclude_key']:
                result[item['exclude_key']] = item['data']
        for item in entity['changed']:
            if item['exclude_key']:
                result[item['exclude_key']] = item['new']
    return result


def resolve_excluded_pairs(excluded_keys, excludable_data):
    """Convierte una lista de exclude_key ('app.Modelo:pk') en pares
    (modelo, data_dict) usando los datos capturados en la última vista
    previa (excludable_data), para reconstruir el predicado de exclusión."""
    from django.apps import apps

    pairs = []
    for key in excluded_keys or []:
        data = excludable_data.get(key)
        if not data:
            continue
        try:
            label, _pk = key.rsplit(':', 1)
            model = apps.get_model(label)
        except (ValueError, LookupError):
            continue
        if model not in EXCLUSION_BUILDERS:
            continue
        pairs.append((model, data))
    return pairs


def cleanup_temp(token):
    if not token:
        return
    shutil.rmtree(temp_dir(token), ignore_errors=True)


def _snapshot_models():
    # Reusa exactamente el mismo formato que ya usa el sistema de auditoría
    # (core/crud_base.py) para que valor_anterior/valor_nuevo sean consistentes.
    from core.crud_base import _model_to_dict
    snapshot = {}
    for model, _label in TRACKED_MODELS:
        snapshot[model] = {obj.pk: _model_to_dict(obj) for obj in model.objects.all()}
    return snapshot


def _display_row(model, data):
    fields = _DISPLAY_FIELDS.get(model, ())
    parts = [str(data.get(f)) for f in fields if data.get(f) not in (None, '')]
    return ' · '.join(parts) if parts else f'#{data.get("pk", "?")}'


def diff_snapshots(before, after):
    """dict: modelo -> {label, new: [...], changed: [...]}"""
    diff = {}
    for model, label in TRACKED_MODELS:
        before_rows = before.get(model, {})
        after_rows = after.get(model, {})
        before_keys = set(before_rows)
        after_keys = set(after_rows)

        excludable = model in EXCLUSION_BUILDERS

        new_items = []
        for pk in sorted(after_keys - before_keys):
            data = after_rows[pk]
            new_items.append({
                'pk': pk, 'data': data, 'display': _display_row(model, data),
                'exclude_key': exclude_key(model, pk) if excludable else None,
            })

        changed_items = []
        for pk in sorted(before_keys & after_keys):
            old, new = before_rows[pk], after_rows[pk]
            changed_fields = {
                k: (old.get(k), new.get(k)) for k in new if new.get(k) != old.get(k)
            }
            if changed_fields:
                changed_items.append({
                    'pk': pk,
                    'old': old,
                    'new': new,
                    'changed_fields': changed_fields,
                    'display': _display_row(model, new),
                    'exclude_key': exclude_key(model, pk) if excludable else None,
                })

        # El importador tambien retira filas (p.ej. asignaciones que dejan de
        # cumplir la validacion de afinidad); deben verse en la vista previa
        # y auditarse igual que una creacion o modificacion. No se ofrece
        # excluir estas: "excluir un retiro" significaria forzar que la fila
        # se quede pese a no cumplir una validacion del importador.
        removed_items = []
        for pk in sorted(before_keys - after_keys):
            data = before_rows[pk]
            removed_items.append({'pk': pk, 'data': data, 'display': _display_row(model, data)})

        if new_items or changed_items or removed_items:
            diff[model] = {
                'label': label, 'new': new_items, 'changed': changed_items,
                'removed': removed_items, 'excludable': excludable,
            }
    return diff


def run_import(base_dir, periodo_codigo, periodo_nombre, *, commit):
    """Ejecuta import_complete_fcacc. Si commit=False calcula el diff dentro
    de una transacción que siempre se revierte (vista previa, no persiste
    nada). Si commit=True, aplica los cambios de verdad.

    Devuelve (diff, steps_results). Puede propagar CommandError si falta un
    archivo obligatorio o el periodo no admite importaciones.
    """
    options = {
        'base_dir': str(base_dir),
        'periodo_codigo': periodo_codigo,
        'periodo_nombre': periodo_nombre,
        'dry_run': False,
        'sync_asignaciones': False,
        'step': 0,
    }

    if commit:
        before = _snapshot_models()
        cmd = ImportCommand(stdout=io.StringIO(), no_color=True)
        cmd.handle(**options)
        after = _snapshot_models()
        steps_results = getattr(cmd, 'last_run_results', {})
        return diff_snapshots(before, after), steps_results

    before = _snapshot_models()
    after = before
    steps_results = {}
    try:
        with transaction.atomic():
            cmd = ImportCommand(stdout=io.StringIO(), no_color=True)
            cmd.handle(**options)
            after = _snapshot_models()
            steps_results = getattr(cmd, 'last_run_results', {})
            raise _PreviewRollback()
    except _PreviewRollback:
        pass
    return diff_snapshots(before, after), steps_results


def apply_audit_log(request, diff):
    """Registra en auditoria_registro_cambios cada fila creada/modificada,
    reusando la misma función que usan las vistas CRUD genéricas."""
    from core.crud_base import _audit_log

    for model, changes in diff.items():
        for item in changes['new']:
            try:
                instance = model.objects.get(pk=item['pk'])
            except model.DoesNotExist:
                continue
            _audit_log(request, instance, 'INSERT')
        for item in changes['changed']:
            try:
                instance = model.objects.get(pk=item['pk'])
            except model.DoesNotExist:
                continue
            _audit_log(request, instance, 'UPDATE', old_values=item['old'])
        for item in changes.get('removed', []):
            # La fila ya no existe en la BD (el importador la retiró, p.ej.
            # por afinidad invalida); se arma una instancia transitoria solo
            # para que _audit_log pueda leer su tabla/pk, sin consultarla.
            instance = model(pk=item['pk'])
            _audit_log(request, instance, 'DELETE', old_values=item['data'], registro_pk=item['pk'])
